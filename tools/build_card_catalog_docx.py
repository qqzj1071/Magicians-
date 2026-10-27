"""Build the v1.0 card catalogue document distributed with the Windows installer."""
from pathlib import Path
import json
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "releases" / "v1.0" / "cards.json"
OUTPUT = ROOT / "releases" / "v1.0" / "Magicians-v1.0-card-catalogue.docx"
ELEMENTS = {"fire": "火", "water": "水", "earth": "土", "wind": "风", "poison": "毒", "generic": "通用"}
RARITIES = {"specialization": "专精", "legendary": "传奇", "epic": "史诗", "rare": "稀有", "common": "普通"}

def main() -> None:
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)
    section.top_margin = Inches(0.35)
    section.bottom_margin = Inches(0.35)
    title = document.add_heading("Magicians! v1.0 卡牌表", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("由 releases/v1.0/cards.json 生成，包含全部卡牌的卡面信息。")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = document.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["卡牌", "元素", "稀有度", "费用", "攻", "防", "首层", "卡面效果"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    cards = sorted(data["cards"].values(), key=lambda c: (c["element"], -int(c["rarity"] == "specialization"), c["rarity"], c["name"]))
    for card in cards:
        cells = table.add_row().cells
        values = [card["name"], ELEMENTS.get(card["element"], card["element"]),
                  RARITIES.get(card["rarity"], card["rarity"]), str(card["cost"]),
                  str(card["attack"]), str(card["defense"]), str(card["layers"]), card["text"]]
        for cell, value in zip(cells, values):
            cell.text = str(value).replace("\n", " ")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
    document.save(OUTPUT)
    print(f"CARD_CATALOGUE_DOCX_OK cards={len(cards)} output={OUTPUT}")

if __name__ == "__main__":
    main()

