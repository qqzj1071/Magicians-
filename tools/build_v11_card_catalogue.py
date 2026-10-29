"""Generate the playtest reference from the current frozen v1.1 card faces."""
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ELEMENTS = {'fire':'火','water':'水','earth':'土','wind':'风','poison':'毒','generic':'通用'}
RARITIES = {'common':'普通','rare':'稀有','epic':'史诗','legendary':'传奇','specialization':'专精'}

def main():
    source = ROOT / 'releases/v1.1/cards.json'
    cards = json.loads(source.read_text(encoding='utf-8'))['cards']
    profiles = {p['CardId']: p for f in (ROOT/'Data/Presentation/Cards').glob('*.json')
                for p in [json.loads(f.read_text(encoding='utf-8-sig'))]}
    assert len(cards) == len(profiles) == 252 and set(cards) == set(profiles)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = section.right_margin = Inches(.55)
    section.top_margin = section.bottom_margin = Inches(.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(9)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(2)
    title = doc.add_paragraph('Magicians v1.1 卡牌表', 'Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0,0,0)
    doc.add_paragraph('本表供 v1.1 试玩查询，收录全部 252 张现行卡牌。卡牌效果和攻防数值来自冻结卡牌库及客户端卡面档案；构筑规则沿用 v1.0。')
    doc.add_paragraph('初始战斗属性：生命 1000、攻击 100、防御 40；基础飞弹间隔 50。费用/研习分别表示战斗费用和基础购买研习点；本轮免费资格按游戏实际提示。首层表示该实体卡首次打出的元素叠层。')
    for element, label in ELEMENTS.items():
        names = sorted((n for n,c in cards.items() if c['element']==element),
                       key=lambda n:(list(RARITIES).index(cards[n]['rarity']),n))
        heading=doc.add_heading(f'{label}系卡牌  {len(names)} 张', 1)
        if element != 'fire':heading.paragraph_format.page_break_before=True
        table=doc.add_table(rows=1, cols=8)
        table.style='Table Grid'
        table.autofit=False
        widths=[1.15,.45,.55,.7,.4,.4,.45,6.44]
        for col,width in zip(table.columns,widths):col.width=Inches(width)
        for cell,text in zip(table.rows[0].cells,['名称','元素','稀有度','费/研习','攻','防','首层','卡面效果']):cell.text=text
        repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for name in names:
            c,p=cards[name],profiles[name]
            assert (p['Attack'],p['Defense'],p['CombatCost'],p['RulesText'])==(c['attack'],c['defense'],c['cost'],c['text'])
            text=p['DisplayRulesText']
            parts=text.split('；')
            timings={'立即','整场','下一发','抽牌前','下抽','下一抽牌阶段前','本次','下一次'}
            while len(parts)>1 and all(t.strip() in timings for t in parts[-1].split('/')):parts.pop()
            text='；'.join(parts).replace('毒附魔层数 +1；','').replace('毒附魔层数+1；','')
            values=[name,label,RARITIES[c['rarity']],f"{c['cost']} / {p['BaseStudyCost']}",str(c['attack']),str(c['defense']),str(c['layers']),text]
            row=table.add_row()
            for cell,value,width in zip(row.cells,values,widths):
                cell.width=Inches(width);cell.text=value
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for row_index,row in enumerate(table.rows):
            for col,cell in enumerate(row.cells):
                for paragraph in cell.paragraphs:
                    if col in range(1,7):paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size=Pt(9)
                        run.bold=row_index==0
    doc.core_properties.title='Magicians v1.1 卡牌表'
    doc.core_properties.author='Magicians'
    doc.core_properties.subject='252 张冻结卡牌与客户端卡面信息'
    for element in [doc.element,doc.styles.element]:
        for border in element.xpath('.//w:pBdr'):border.getparent().remove(border)
    output=ROOT/'releases/v1.1/Magicians-v1.1-card-catalogue.docx'
    doc.save(output)
    print(json.dumps(dict(cards=len(cards),output=str(output),cards_sha256=hashlib.sha256(source.read_bytes()).hexdigest()),ensure_ascii=False))

if __name__=='__main__':main()
